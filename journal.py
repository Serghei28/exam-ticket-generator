import os
import random
import tempfile
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook, load_workbook


HEADERS = ("Last name", "First name", "Номер билета", "Дата и время")
DATE_FORMAT = "dd.mm.yyyy hh:mm:ss"


class JournalError(Exception):
    """Ошибка структуры или чтения журнала."""


def generate_ticket() -> int:
    """Возвращает случайный номер билета от 1 до 20."""
    return random.randint(1, 20)


def _load_or_create_workbook(path: Path):
    if not path.exists():
        workbook = Workbook()
        if not workbook.worksheets:
            workbook.close()
            raise JournalError(f"В {path.name} нет листов с журналом.")

        worksheet = workbook.active
        worksheet.append(HEADERS)
        return workbook, worksheet

    try:
        workbook = load_workbook(path)
    except PermissionError:
        raise
    except Exception as error:
        raise JournalError(
            f"Не удалось открыть {path.name}: файл поврежден или имеет неверный формат."
        ) from error

    worksheet = workbook.active
    actual_headers = tuple(worksheet.cell(row=1, column=column).value for column in range(1, 5))
    if actual_headers != HEADERS:
        workbook.close()
        raise JournalError(
            f"В {path.name} неверные заголовки. Ожидались: {', '.join(HEADERS)}."
        )

    return workbook, worksheet


def append_record(
    path: Path,
    last_name: str,
    first_name: str,
    ticket_number: int,
    timestamp: datetime,
) -> None:
    """Добавляет запись и безопасно заменяет файл журнала после успешного сохранения."""
    workbook, worksheet = _load_or_create_workbook(path)
    temporary_path = None

    try:
        worksheet.append((last_name, first_name, ticket_number, timestamp))
        worksheet.cell(row=worksheet.max_row, column=4).number_format = DATE_FORMAT

        with tempfile.NamedTemporaryFile(
            mode="wb",
            prefix=".journal-",
            suffix=".xlsx",
            dir=path.parent,
            delete=False,
        ) as temporary_file:
            temporary_path = Path(temporary_file.name)

        workbook.save(temporary_path)
        workbook.close()
        os.replace(temporary_path, path)
        temporary_path = None
    finally:
        try:
            workbook.close()
        finally:
            if temporary_path is not None:
                try:
                    temporary_path.unlink()
                except FileNotFoundError:
                    pass
