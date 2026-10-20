import tempfile
import unittest
from datetime import datetime
from pathlib import Path
from unittest.mock import patch

from openpyxl import load_workbook

import journal


class JournalTests(unittest.TestCase):
    def test_creates_workbook_with_headers_and_record(self):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "journal.xlsx"
            timestamp = datetime(2026, 10, 20, 22, 0, 0)

            journal.append_record(path, "Иванов", "Иван", 7, timestamp)

            workbook = load_workbook(path)
            worksheet = workbook.active
            self.assertEqual(
                tuple(cell.value for cell in worksheet[1][:4]),
                journal.HEADERS,
            )
            self.assertEqual(
                tuple(cell.value for cell in worksheet[2][:4]),
                ("Иванов", "Иван", 7, timestamp),
            )
            self.assertEqual(worksheet["D2"].number_format, journal.DATE_FORMAT)
            workbook.close()

    def test_appends_without_changing_previous_rows(self):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "journal.xlsx"
            first_timestamp = datetime(2026, 10, 20, 22, 1, 0)
            second_timestamp = datetime(2026, 10, 20, 22, 2, 0)
            journal.append_record(path, "Первый", "Студент", 3, first_timestamp)
            first_row_before = ("Первый", "Студент", 3, first_timestamp)

            journal.append_record(path, "Второй", "Студент", 19, second_timestamp)

            workbook = load_workbook(path)
            worksheet = workbook.active
            self.assertEqual(
                tuple(cell.value for cell in worksheet[2][:4]),
                first_row_before,
            )
            self.assertEqual(
                tuple(cell.value for cell in worksheet[3][:4]),
                ("Второй", "Студент", 19, second_timestamp),
            )
            workbook.close()

    def test_failed_replace_cleans_temporary_file_and_retry_has_one_record(self):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "journal.xlsx"
            real_replace = journal.os.replace
            replace_calls = 0

            def replace_once_then_succeed(source, target):
                nonlocal replace_calls
                replace_calls += 1
                if replace_calls == 1:
                    raise PermissionError("locked")
                return real_replace(source, target)

            with patch("journal.os.replace", side_effect=replace_once_then_succeed):
                with self.assertRaises(PermissionError):
                    journal.append_record(
                        path,
                        "Иванов",
                        "Иван",
                        7,
                        datetime(2026, 10, 20, 22, 3, 0),
                    )
                journal.append_record(
                    path,
                    "Иванов",
                    "Иван",
                    7,
                    datetime(2026, 10, 20, 22, 3, 0),
                )

            workbook = load_workbook(path)
            self.assertEqual(workbook.active.max_row, 2)
            workbook.close()
            self.assertEqual(list(Path(directory).glob(".journal-*.xlsx")), [])

    def test_generate_ticket_is_always_in_range(self):
        tickets = [journal.generate_ticket() for _ in range(200)]

        self.assertTrue(all(1 <= ticket <= 20 for ticket in tickets))


if __name__ == "__main__":
    unittest.main()
